"""Fill the official SIH 2026 IDEA Presentation template for problem statement SIH26043.

Outputs backend/SIH_Presentation_PS26043.pptx (6 slides, template format preserved).
"""
import shutil
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Pt

ROOT = Path(__file__).resolve().parent
TEMPLATE_SRC = Path(sys.argv[1] if len(sys.argv) > 1 else "/tmp/SIH2026-IDEA-Presentation-Format.pptx")
TEMPLATE_LOCAL = ROOT / "SIH2026-IDEA-Presentation-Format.pptx"
OUT = ROOT / "SIH_Presentation_PS26043.pptx"

shutil.copy(TEMPLATE_SRC, TEMPLATE_LOCAL)
prs = Presentation(TEMPLATE_LOCAL)

FONT = "Arial"


def set_lines(shape, lines, base=12):
    """Replace a shape's text with bulleted lines.

    lines: list of tuples (level, text). Returns the text frame.
    """
    tf = shape.text_frame
    for para in list(tf.paragraphs)[1:]:
        para._p.getparent().remove(para._p)
    first = tf.paragraphs[0]
    for run in list(first.runs):
        run._r.getparent().remove(run._r)
    first.level = 0
    for i, (level, text) in enumerate(lines):
        para = first if i == 0 else tf.add_paragraph()
        para.level = level
        para.space_after = Pt(3)
        para.space_before = Pt(0)
        run = para.add_run()
        run.text = text
        run.font.name = FONT
        run.font.bold = level == 0
        run.font.size = Pt(base if level == 0 else base - 1.5)
    return tf


def rename_team_ovals(team):
    for slide in prs.slides:
        for sh in slide.shapes:
            if sh.name.startswith("Oval") and sh.has_text_frame and "Your Team Name" in sh.text_frame.text:
                set_lines(sh, [(0, team)], base=11)


def title_page():
    slide = prs.slides[0]
    for sh in slide.shapes:
        if sh.name == "TextBox 9":
            lines = [
                (0, "Problem Statement ID – SIH26043"),
                (0, "Problem Statement Title – Digital platform to crowdsource societal challenges and facilitate collaborative problem solving through universities and industry partnerships"),
                (0, "Theme – Smart Education"),
                (0, "PS Category – Software"),
                (0, "Team ID – TBD"),
                (0, "Team Name (Registered on portal) – Pheonix Coders"),
            ]
            tf = set_lines(sh, lines, base=18)
            for para in tf.paragraphs:
                para.space_after = Pt(6)
        if sh.name == "Subtitle 3":
            sh.text_frame.text = "Setu — AI-Augmented Societal Problem-Solving Ecosystem"


def content_slide(idx, title, lines):
    slide = prs.slides[idx]
    for sh in slide.shapes:
        if sh.name == "Title 1":
            sh.text_frame.text = title
        elif sh.name == "TextBox 8":
            set_lines(sh, lines, base=14)


# Slide 1 — title page
title_page()
rename_team_ovals("Pheonix Coders")

# Slide 2 — Proposed Solution
content_slide(
    1,
    "Setu: A Citizen-to-Solution Innovation Pipeline",
    [
        (0, "One transparent pipeline — from a citizen-reported problem to a verified, deployed solution."),
        (1, "Citizen reporting (no login): web form + photos/video + GPS into a structured queue"),
        (1, "AI Articulation Wizard: raw input → structured, de-duplicated challenge statements (LLM + offline rule-based fallback)"),
        (1, "Semantic Matching Engine: sentence-transformer embeddings + tag overlap route each challenge to the right university teams"),
        (1, "Lifecycle & Workflow Tracker: review → team formation → milestones → deployment → outcome verification with evidence"),
        (1, "Industry Collaboration Hub: co-develop, fund, mentor; IP, patents and startups tracked"),
        (1, "Govt Analytics Dashboard: domain, district, institution, industry, community, patent and startup metrics"),
        (0, "How it addresses the problem"),
        (1, "Adds the missing centralized channel to collect, categorize and route societal challenges"),
        (1, "Replaces fragmented university–industry collaboration with a transparent ecosystem"),
        (0, "Innovation & uniqueness"),
        (1, "Persistent lifecycle and verified outcomes — a project is not 'done' until deployed and confirmed on ground"),
        (1, "Offline-first AI, low infrastructure cost, role-based journeys for 6 personas"),
    ],
)

# Slide 3 — Technical Approach
content_slide(
    2,
    "TECHNICAL APPROACH",
    [
        (0, "Technology stack"),
        (1, "Frontend: React 18 + Vite + TailwindCSS + ECharts — fast, accessible, mobile-first UI"),
        (1, "Backend: FastAPI (Python) REST API, OAuth2/JWT, role-based access control (6 personas)"),
        (1, "AI: Gemini Flash (Wizard) + sentence-transformers paraphrase-multilingual-MiniLM-L12-v2 (matching) + rule-based fallback"),
        (1, "Data: SQLite (migratable to PostgreSQL); embeddings as JSON, cosine nearest-neighbour search"),
        (0, "Methodology & implementation"),
        (1, "5-step AI wizard → publish → de-duplicate → semantic match → interest → team → milestones → verify"),
        (1, "Working prototype live and demoable for every role: setu-mishtiiikrishnaa19.vercel.app"),
        (0, "Architecture flow"),
        (1, "Citizen/Dept input → Wizard → Marketplace → Match → Project → Milestones → Verify → Dashboard"),
    ],
)

# Slide 4 — Feasibility and Viability
content_slide(
    3,
    "FEASIBILITY AND VIABILITY",
    [
        (0, "Technical feasibility"),
        (1, "Proven open-source stack; low compute — embeddings run on CPU"),
        (1, "Graceful degradation: rule-based fallback keeps the platform working without LLM APIs"),
        (0, "Potential challenges & mitigation"),
        (1, "Cold-start / adoption → realistic seeded Jharkhand dataset; citizen-first reporting lowers the entry barrier"),
        (1, "LLM cost & latency → caching plus offline fallback"),
        (1, "Fair matching & verification bias → transparent match scores; evidence + problem-owner confirmation"),
        (0, "Operational & financial viability"),
        (1, "Lightweight onboarding via AI wizard; workflow mirrors department SOPs"),
        (1, "No commercial vector DB; cheap hosting; ROI measured via verified on-ground outcomes"),
    ],
)

# Slide 5 — Impact and Benefits
content_slide(
    4,
    "IMPACT AND BENEFITS",
    [
        (0, "Impact on target audience"),
        (1, "Government: measurable social outcomes, transparent dashboards, faster problem-to-research pipeline"),
        (1, "Universities & students: real-world experiential learning (NEP 2020), funding, IP and startup pathways"),
        (1, "Industry, startups, MSMEs & CSR: demand-driven opportunities, talent pipeline, measurable CSR impact"),
        (1, "Citizens: local issues enter an accountable, trackable pipeline"),
        (0, "Benefits of the solution"),
        (1, "Social: civic problems are solved instead of shelved"),
        (1, "Economic: startups and patents born from real community needs"),
        (1, "Institutional: multidisciplinary, demand-driven research collaboration"),
    ],
)

# Slide 6 — Research and References
content_slide(
    5,
    "RESEARCH AND REFERENCES",
    [
        (1, "SIH 2026 problem statement SIH26043 — sih.gov.in/sih2026PS (Govt of Jharkhand, Dept of Higher & Technical Education)"),
        (1, "National Education Policy 2020 — experiential learning, community engagement, industry collaboration"),
        (1, "Jharkhand Student Research and Innovation Policy (2024)"),
        (1, "Gap analysis of existing platforms (SIH portal, YUKTI, Kavach): no persistent lifecycle, matching or verification"),
        (1, "Sentence-Transformers (SBERT); React, Vite, FastAPI official documentation"),
        (1, "Prototype: github.com/mishtiiikrishnaa/sih2026prototype — live: setu-mishtiiikrishnaa19.vercel.app"),
    ],
)

# Drop the 7th slide (instructions) — SIH allows a maximum of 6 slides including the title.
slide_id_lst = prs.slides._sldIdLst
slides = list(slide_id_lst)
rId = slides[-1].rId
slide_id_lst.remove(slides[-1])
prs.part.drop_rel(rId)

prs.save(OUT)
print(f"Saved {OUT}: {len(prs.slides._sldIdLst)} slides")