import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

def add_slide(prs, title_text, content_text):
    slide_layout = prs.slide_layouts[1] # Title and Content
    slide = prs.slides.add_slide(slide_layout)
    
    title = slide.shapes.title
    title.text = title_text
    
    # Optional styling for title
    for paragraph in title.text_frame.paragraphs:
        for run in paragraph.runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(0, 51, 102)
    
    body = slide.placeholders[1]
    body.text = content_text
    
    # Optional styling for body
    for paragraph in body.text_frame.paragraphs:
        paragraph.font.size = Pt(16)
        if paragraph.text.startswith("-"):
            paragraph.level = 0
        elif paragraph.text.startswith("  -"):
            paragraph.level = 1

prs = Presentation()

# Slide 1: Title
title_slide_layout = prs.slide_layouts[0]
slide = prs.slides.add_slide(title_slide_layout)
title = slide.shapes.title
subtitle = slide.placeholders[1]
title.text = "Setu: AI-Augmented Societal Problem-Solving Ecosystem"
subtitle.text = "Smart India Hackathon 2026\nProblem Statement: PS26043\nTeam: Team Antigravity\nGovernment of Jharkhand"

# Slide 2: Proposed Solution
add_slide(prs, "Proposed Solution: Setu",
"""- A tri-sided marketplace connecting Government (Problem Owners), Universities (Solvers), and Administrators.
- Core Gaps Addressed:
  - No continuous lifecycle tracking (hackathons end abruptly)
  - Poor problem articulation from domain experts
  - Lack of semantic matching between solver expertise and problems
- Key Features:
  - AI Articulation Wizard: Translates raw situations into structured challenges using LLMs.
  - Semantic Matching Engine: Vector-based matching (Sentence-Transformers) + Tag overlap.
  - Lifecycle Tracker: End-to-end tracking from Discovery to Deployment and Outcome Verification.""")

# Slide 3: Technical Approach
add_slide(prs, "Technical Approach & Architecture",
"""- Frontend: React + Vite + TailwindCSS (Fast, accessible, mobile-responsive UI).
- Backend: FastAPI (Python) for high performance and asynchronous execution.
- AI/ML Layer: 
  - Sentence-Transformers (paraphrase-multilingual-MiniLM-L12-v2) for semantic matching.
  - Gemini Flash API for the AI Articulation Wizard (with rule-based NLP fallback).
- Database: SQLite (Scalable to PostgreSQL). Embeddings stored as JSON.
- Design: State-driven role-based access control (RBAC) across 4 personas.""")

# Slide 4: Feasibility & Viability
add_slide(prs, "Feasibility & Viability",
"""- Technical Feasibility: 
  - Built on open-source, proven technologies (React, FastAPI, SentenceTransformers).
  - Designed for offline/low-resource fallback (rule-based matching).
- Operational Viability:
  - Lightweight onboarding for Government officials via AI Wizard.
  - Minimal infrastructure cost (CPU-runnable embedding models).
- Financial Viability:
  - Avoids expensive commercial vector DBs; uses in-memory cosine similarity and SQL JSON.
  - High ROI for government by verifying actual on-ground deployment.""")

# Slide 5: Impact & Benefits
add_slide(prs, "Impact & Benefits",
"""- For Government:
  - Transparent dashboard for state-level innovation metrics.
  - Faster translation of citizen issues into academic research.
- For Universities/Students:
  - Direct pipeline to real-world impact and funding.
  - Skill-based matching ensures relevant project allocation.
- For Citizens:
  - Verified outcomes (Project isn't "done" until deployment is validated).
  - Scalable model across domains: Water, Healthcare, Education.""")

# Slide 6: Research & References
add_slide(prs, "Research, References & Future Scope",
"""- Extensive R&D conducted on existing platforms (Kavach, SIH, YUKTI) showing absence of persistent lifecycle tracking.
- Future Roadmap:
  - Integration with DigiLocker for student identity verification.
  - Multilingual AI Wizard (Hindi/Santali voice input) for rural stakeholders.
  - Blockchain-based verification for project milestones and funding disbursement.
- References: 
  - Jharkhand Student Research and Innovation Policy (2024).
  - Sentence-Transformers Documentation.""")

prs.save("SIH_Presentation.pptx")
print("Presentation saved as SIH_Presentation.pptx")
